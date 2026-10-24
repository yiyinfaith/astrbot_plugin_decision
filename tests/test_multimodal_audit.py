"""Cross-session isolation, cache sufficiency and response boundary regressions."""

from __future__ import annotations

import asyncio
import time
from types import SimpleNamespace

import pytest

from tests.personal_network._astrbot_stub import install

install()

from decision.multimodal.adapters import _extract_text  # noqa: E402
from decision.multimodal.cache import MultimodalCache  # noqa: E402
from decision.multimodal.models import MediaIndexEntry, MediaRef, Understanding  # noqa: E402
from tests.test_multimodal_runtime import FakeProvider, build, png  # noqa: E402


@pytest.mark.parametrize("same_session", [False, True])
async def test_inflight_respects_session_and_reuses_generic_wording(tmp_path, same_session):
    provider = FakeProvider()
    started, release = asyncio.Event(), asyncio.Event()

    async def blocked(**kwargs):
        provider.calls.append(kwargs)
        number = len(provider.calls)
        started.set()
        await release.wait()
        return SimpleNamespace(completion_text=f"真实调用 {number} 的画面描述")

    provider.text_chat = blocked
    runtime = build(tmp_path, provider)
    path = png(tmp_path)
    first = MediaRef(kind="image", reference=str(path), urn="first", umo="session-one")
    second = MediaRef(kind="image", reference=str(path), urn="second", umo="session-one" if same_session else "session-two")
    tasks = []
    try:
        tasks.append(asyncio.create_task(runtime.understand(first)))
        await asyncio.wait_for(started.wait(), 1)
        tasks.append(asyncio.create_task(runtime.understand(second, instruction="请描述画面" if same_session else "")))
        await asyncio.sleep(0.05)
        assert len(provider.calls) == (1 if same_session else 2)
        release.set()
        one, two = await asyncio.gather(*tasks)
        assert one.umo == first.umo and two.umo == second.umo
        if not same_session:
            assert one.key != two.key and one.summary != two.summary
    finally:
        release.set()
        await asyncio.gather(*tasks, return_exceptions=True)
        await runtime.terminate()


@pytest.mark.parametrize("question", ["车牌号是什么", "Read the serial number", "Translate the sign into French"])
async def test_specific_short_followup_is_an_incremental_read(tmp_path, question):
    provider = FakeProvider()
    runtime = build(tmp_path, provider)
    ref = MediaRef(kind="image", reference=str(png(tmp_path)), urn="first", umo="session")
    try:
        initial = await runtime.understand(ref)
        followup = await runtime.understand(ref, instruction=question)
        assert initial.ok and followup.ok
        assert len(provider.calls) == 2
        assert initial.summary in provider.calls[1]["prompt"]
        assert question in provider.calls[1]["prompt"]
        assert initial.summary in followup.summary
        await runtime.understand(ref, instruction=question)
        assert len(provider.calls) == 2
    finally:
        await runtime.terminate()


@pytest.mark.parametrize("response", [
    SimpleNamespace(completion_text="", reasoning_content="PRIVATE_THOUGHT", api_key="PRIVATE_KEY"),
    {"reasoning_content": "PRIVATE_THOUGHT", "tool_calls": ["not an answer"]},
    object(),
])
def test_only_final_provider_text_is_a_successful_answer(response):
    assert _extract_text(response) == ""


async def test_clearing_index_only_session_is_durable(tmp_path):
    path = tmp_path / "cache.json"
    cache = MultimodalCache(path)
    await cache.remember_media([MediaIndexEntry(
        umo="session", message_id="one", kind="image", index=0, name="private.png",
        reference="/private.png", content_key="media", received_at=time.time(),
    )])
    assert await cache.recent_media("session")
    await cache.clear_session("session")
    assert await MultimodalCache(path).recent_media("session") == []


async def test_malformed_index_payload_is_ignored_without_losing_valid_entries(tmp_path):
    path = tmp_path / "cache.json"
    path.write_text('{"entries": [], "index": ["invalid"]}', encoding="utf-8")
    cache = MultimodalCache(path)
    await cache.load()
    assert await cache.recent_media("session") == []


async def test_concurrent_load_cannot_overwrite_index_before_initialization(tmp_path, monkeypatch):
    cache = MultimodalCache(tmp_path / "cache.json")
    started, release = asyncio.Event(), asyncio.Event()
    original = asyncio.to_thread

    async def blocked_load(function, *args, **kwargs):
        if getattr(function, "__name__", "") == "read_text":
            started.set()
            await release.wait()
        return await original(function, *args, **kwargs)

    monkeypatch.setattr(asyncio, "to_thread", blocked_load)
    loading = asyncio.create_task(cache.load())
    writing = None
    try:
        await started.wait()
        writing = asyncio.create_task(cache.put(Understanding(key="new", kind="image", umo="s", summary="new")))
        await asyncio.sleep(0.05)
        assert not writing.done(), "a writer must wait until the on-disk cache finishes loading"
    finally:
        release.set()
        await loading
        if writing is not None:
            await writing


async def test_long_video_focus_requests_preserve_user_question_and_frame_budget(tmp_path, monkeypatch):
    from decision.multimodal import adapters, preprocess

    provider = FakeProvider()
    runtime = build(tmp_path, provider, mm_video_max_frames=2, mm_video_agentic_rounds=1)

    async def probe(path):
        return {"duration": 90, "has_audio": False}

    frame_counts = []

    async def frames(path, workdir, stamps, **kwargs):
        frame_counts.append(len(stamps))
        return [str(png(tmp_path)) for _ in stamps]

    prompts = []

    async def call(provider, **kwargs):
        prompts.append(kwargs["prompt"])
        return "10-15: 有关序列号的关键片段"

    monkeypatch.setattr(preprocess, "probe_video", probe)
    monkeypatch.setattr(preprocess, "extract_frames", frames)
    monkeypatch.setattr(adapters, "call_provider", call)
    path = tmp_path / "video.mp4"
    path.write_bytes(b"test")
    try:
        result = await runtime.understand(MediaRef(kind="video", reference=str(path), umo="s", urn="video"), instruction="只核对设备序列号 SERIAL-7429")
        assert result.ok
        assert prompts and all("只核对设备序列号 SERIAL-7429" in prompt for prompt in prompts)
        assert max(frame_counts) <= 2
    finally:
        await runtime.terminate()


@pytest.mark.parametrize("address", ["224.0.0.1", "239.255.255.250", "ff02::1", "ff0e::1", "fec0::1", "64:ff9b::a00:1"])
def test_ssrf_rejects_special_global_addresses(address):
    from decision.multimodal.media import _is_public_ip

    assert _is_public_ip(address) is False


def test_idn_resolution_uses_the_same_ascii_host_as_https_sni(monkeypatch):
    import socket

    from decision.multimodal.media import _public_resolution

    monkeypatch.setattr(socket, "getaddrinfo", lambda *args, **kwargs: [(socket.AF_INET, socket.SOCK_STREAM, 6, "", ("93.184.216.34", 443))])
    host, port, pinned = _public_resolution("https://例子.com/image.png")
    assert host == "例子.com".encode("idna").decode("ascii")
    assert pinned[0]["hostname"] == host and port == 443


def test_pptx_ordinal_target_is_not_confused_by_digits_in_other_titles(tmp_path):
    from pptx import Presentation

    from decision.multimodal.office import parse_document, read_target

    book = Presentation()
    for title in ["2026 report -- PRIVATE_FIRST", "ONLY_SECOND"]:
        book.slides.add_slide(book.slide_layouts[0]).shapes.title.text = title
    path = tmp_path / "slides.pptx"
    book.save(path)
    result = read_target(parse_document(path), "2")
    assert "ONLY_SECOND" in result and "PRIVATE_FIRST" not in result


def test_pdf_target_range_is_lazy_and_rejects_invalid_selectors():
    from decision.multimodal.runtime import MultimodalRuntime

    selected = MultimodalRuntime._target_pdf_pages("1-999999999999")
    assert isinstance(selected, range) and selected[0] == 0
    assert 123456 in selected
    assert list(MultimodalRuntime._target_pdf_pages("page 2")) == [1]
    for target in ("nonsense", "0", "5-2"):
        with pytest.raises(ValueError):
            MultimodalRuntime._target_pdf_pages(target)


def test_pdf_first_page_longer_than_budget_is_not_silently_empty(tmp_path):
    pytest.importorskip("pypdf")
    import fitz

    from decision.multimodal.preprocess import pdf_text_pages

    path = tmp_path / "long.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_textbox((20, 20, 550, 600), "UNIQUE-FIRST-LINE " + "native text " * 80)
    doc.save(path)
    doc.close()
    result = pdf_text_pages(path, 80, range(0, 999999999999))
    assert "UNIQUE-FIRST-LINE" in result and len(result) <= 80


def test_native_generate_content_excludes_thought_parts():
    from decision.multimodal.native_video import _generate_content_text, _safe_error

    payload = {"candidates": [{"content": {"parts": [
        {"thought": True, "text": "PRIVATE_THOUGHT"},
        {"text": "FINAL ANSWER"},
    ]}}]}
    assert _generate_content_text(payload) == "FINAL ANSWER"
    secret = "secret-" + "x" * 300
    assert "secret-" not in _safe_error(RuntimeError("error " + secret), secret)


async def test_provider_failure_is_redacted_before_cache_and_public_tool(tmp_path):
    provider = FakeProvider()
    secret = "super-private-" + "x" * 250
    provider.provider_config["key"] = [secret]

    async def fail(**kwargs):
        raise RuntimeError("request rejected: " + secret)

    provider.text_chat = fail
    runtime = build(tmp_path, provider)
    try:
        item = await runtime.understand(MediaRef(kind="image", reference=str(png(tmp_path)), urn="x", umo="s"))
        assert not item.ok and "super-private" not in item.error
        assert "<redacted>" in item.error
        assert "super-private" not in runtime.cache.path.read_text(encoding="utf-8")
    finally:
        await runtime.terminate()


def test_jev_score_rejects_nan():
    from decision.models import DecisionValidationError, parse_systemone_response

    with pytest.raises(DecisionValidationError):
        parse_systemone_response({"answers": {"score": {"score": float("nan")}}}, {"score": {"type": "score", "criteria": ["low", "high"]}})


async def test_inflight_does_not_reuse_a_different_processing_strategy(tmp_path):
    provider = FakeProvider()
    started, release = asyncio.Event(), asyncio.Event()

    async def blocked(**kwargs):
        provider.calls.append(kwargs)
        started.set()
        await release.wait()
        return SimpleNamespace(completion_text="overview")

    provider.text_chat = blocked
    runtime = build(tmp_path, provider)
    ref = MediaRef(kind="image", reference=str(png(tmp_path)), urn="x", umo="s")
    first = asyncio.create_task(runtime.understand(ref))
    second = None
    try:
        await started.wait()
        runtime.refresh({**runtime.settings.values, "mm_image_max_edge": 600})
        second = asyncio.create_task(runtime.understand(ref, instruction="请描述画面"))
        await asyncio.sleep(.05)
        assert len(provider.calls) == 2
    finally:
        release.set()
        await asyncio.gather(*[task for task in (first, second) if task is not None], return_exceptions=True)
        await runtime.terminate()


async def test_silent_video_track_never_requests_audio_model(tmp_path, monkeypatch):
    from decision.multimodal import adapters, preprocess

    provider = FakeProvider()
    provider.provider_config["modalities"] = ["text", "image", "audio"]
    runtime = build(tmp_path, provider)
    audio = tmp_path / "silent.wav"
    audio.write_bytes(b"test")

    async def extract(*args):
        return str(audio)

    async def silent(*args):
        return True

    monkeypatch.setattr(preprocess, "extract_audio_track", extract)
    monkeypatch.setattr(preprocess, "audio_is_silent", silent)
    try:
        choice, error = adapters.pick_provider(runtime._context(), runtime.settings, "video")
        assert not error
        text = await runtime._video_transcript(choice, tmp_path / "video.mp4", tmp_path, 10)
        assert "静音" in text and provider.calls == []
        assert "未验证" in runtime._video_audio_note({"has_audio": True}, "")
        assert "没有音轨" in runtime._video_audio_note({"has_audio": False}, "")
    finally:
        await runtime.terminate()


async def test_cancelling_one_waiter_preserves_shared_call_and_reload_stops_tools(tmp_path):
    provider = FakeProvider()
    started, release = asyncio.Event(), asyncio.Event()

    async def blocked(**kwargs):
        provider.calls.append(kwargs)
        started.set()
        await release.wait()
        return SimpleNamespace(completion_text="shared")

    provider.text_chat = blocked
    runtime = build(tmp_path, provider)
    ref = MediaRef(kind="image", reference=str(png(tmp_path)), urn="x", umo="s")
    first = asyncio.create_task(runtime.understand(ref))
    await started.wait()
    second = asyncio.create_task(runtime.understand(ref, instruction="请描述画面"))
    await asyncio.sleep(.03)
    first.cancel()
    with pytest.raises(asyncio.CancelledError):
        await first
    assert not second.done() and len(provider.calls) == 1
    await runtime.terminate()
    with pytest.raises(asyncio.CancelledError):
        await second
    assert not runtime._tasks and not runtime._inflight
    after = await runtime.understand(ref)
    assert not after.ok and len(provider.calls) == 1


async def test_component_urls_do_not_bypass_safe_downloader_via_framework_converter():
    from decision.multimodal.media import _resolve_reference

    class Image:
        url = "http://127.0.0.1/private"

        async def convert_to_file_path(self):
            raise AssertionError("framework must not make an unvalidated request")

    assert await _resolve_reference(Image()) == Image.url


async def test_extensionless_download_uses_original_attachment_format_and_cleans_up(tmp_path, monkeypatch):
    from openpyxl import Workbook

    from decision.multimodal import media

    source = tmp_path / "download.bin"
    book = Workbook()
    book.active["A401"] = "EXACT-ATTACHMENT-7429"
    book.active["B401"] = "PRIVATE-OTHER-COLUMN"
    book.save(source)
    limits = []

    async def downloaded(reference, **kwargs):
        limits.append(kwargs["max_bytes"])
        return str(source)

    monkeypatch.setattr(media, "ensure_local", downloaded)
    provider = FakeProvider()
    runtime = build(tmp_path, provider, mm_max_file_mb=5)
    try:
        item = await runtime.understand(MediaRef(kind="file", reference="https://example.invalid/download?id=1", name="original.xlsx", umo="s", urn="x", extra={"target": "Sheet!A401"}))
        assert item.ok, item.error
        assert limits == [5 * 1024 * 1024]
        assert "EXACT-ATTACHMENT-7429" in provider.calls[0]["prompt"]
        assert "PRIVATE-OTHER-COLUMN" not in provider.calls[0]["prompt"]
        assert not source.exists()
    finally:
        await runtime.terminate()


def test_office_page_ranges_and_long_first_block_preserve_requested_content():
    from decision.multimodal.office import DocumentIndex, _page_selection

    assert _page_selection("2-3", 5) == [1, 2]
    assert _page_selection("page 2-999999999999", 4) == [1, 2, 3]
    assert _page_selection("1,3", 4) == [0, 2]
    for target in ("0", "4-2", "6", "garbage"):
        with pytest.raises(ValueError):
            _page_selection(target, 5)
    doc = DocumentIndex(kind="docx", blocks=[{"label": "正文", "text": "UNIQUE-FIRST-7429 " + "long " * 500}])
    assert "UNIQUE-FIRST-7429" in doc.render(limit_chars=200)


def test_native_provider_base_and_key_list_match_sdk_configuration():
    from decision.multimodal.native_video import gemini_credentials

    provider = SimpleNamespace(provider_config={"type": "gemini_chat_completion", "key": ["private-key"], "api_base": "https://example.invalid/v1beta/"})
    assert gemini_credentials(provider) == ("private-key", "https://example.invalid")


async def test_multimodal_controls_transport_retry_attempts_without_multiplication():
    from decision.multimodal.adapters import call_provider

    received = []

    async def chat(*, prompt, request_max_retries=None):
        received.append(request_max_retries)
        return SimpleNamespace(completion_text="ok")

    assert await call_provider(SimpleNamespace(text_chat=chat), prompt="audit") == "ok"
    assert received == [1]


def test_empty_final_model_output_never_falls_back_to_an_earlier_intermediate_answer():
    from decision.multimodal.native_video import _model_output_text

    assert _model_output_text({"steps": [
        {"type": "model_output", "content": [{"type": "text", "text": "INTERMEDIATE"}]},
        {"type": "function_call", "name": "inspect"},
        {"type": "model_output", "content": []},
    ]}) == ""
