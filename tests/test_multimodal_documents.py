"""Office deep parsing, local OCR and native video routing."""

from __future__ import annotations

import asyncio
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from tests.personal_network._astrbot_stub import install  # noqa: E402

install()

from decision.multimodal import native_video, ocr, office  # noqa: E402


# ------------------------------------------------------------------- fixtures
@pytest.fixture()
def docx_file(tmp_path):
    import docx

    document = docx.Document()
    document.add_heading("第一章 概述", level=1)
    document.add_paragraph("这是第一章的正文。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "名称"
    table.cell(0, 1).text = "数值"
    table.cell(1, 0).text = "温度"
    table.cell(1, 1).text = "26"
    document.add_heading("第二章 细节", level=2)
    document.add_paragraph("第二章只讲细节。")
    path = tmp_path / "a.docx"
    document.save(path)
    return path


@pytest.fixture()
def pptx_file(tmp_path):
    from pptx import Presentation

    presentation = Presentation()
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "封面标题"
    first.placeholders[1].text = "副标题内容"
    second = presentation.slides.add_slide(presentation.slide_layouts[5])
    second.shapes.title.text = "数据页"
    table = second.shapes.add_table(2, 2, 0, 0, 100, 100).table
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "值"
    second.notes_slide.notes_text_frame.text = "讲者备注"
    path = tmp_path / "a.pptx"
    presentation.save(path)
    return path


@pytest.fixture()
def xlsx_file(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据表"
    sheet.append(["姓名", "分数"])
    sheet.append(["小明", 90])
    sheet["C1"] = "=B2*2"
    sheet.merge_cells("A4:B4")
    workbook.create_sheet("说明").append(["附注"])
    path = tmp_path / "a.xlsx"
    workbook.save(path)
    return path


# --------------------------------------------------------------- Office: DOCX
def test_docx_keeps_headings_tables_and_sections(docx_file):
    index = office.parse_document(docx_file)
    kinds = [block["kind"] for block in index.blocks]
    assert "heading" in kinds and "paragraph" in kinds and "table" in kinds
    assert index.meta["sections"] == 2
    rendered = index.render()
    assert "第一章 概述" in rendered
    assert "温度" in rendered and "26" in rendered


def test_docx_can_be_read_by_section_or_table(docx_file):
    index = office.parse_document(docx_file)
    second = office.read_target(index, "第二章")
    assert "第二章只讲细节" in second
    assert "第一章" not in second, "targeted read must exclude other sections"
    table_text = office.read_target(index, "表格")
    assert "温度" in table_text


def test_unknown_target_lists_available_ranges(docx_file):
    index = office.parse_document(docx_file)
    with pytest.raises(ValueError) as exc:
        office.read_target(index, "不存在的章节")
    assert "可用范围" in str(exc.value)


# --------------------------------------------------------------- Office: PPTX
def test_pptx_reads_slides_tables_and_notes(pptx_file):
    index = office.parse_document(pptx_file)
    assert index.meta["slides"] == 2
    rendered = index.render()
    assert "封面标题" in rendered
    assert "项目" in rendered, "table cells must be captured"
    assert "讲者备注" in rendered, "speaker notes must be captured"
    assert "[幻灯片 2]" in rendered


def test_pptx_single_slide_read(pptx_file):
    index = office.parse_document(pptx_file)
    only_second = office.read_target(index, "2")
    assert "数据页" in only_second
    assert "封面标题" not in only_second


# --------------------------------------------------------------- Office: XLSX
def test_xlsx_keeps_sheets_rows_formulas_and_merges(xlsx_file):
    index = office.parse_document(xlsx_file)
    assert index.meta["sheets"] == ["数据表", "说明"]
    rendered = index.render()
    assert "[工作表 数据表]" in rendered
    assert "第 2 行" in rendered and "小明" in rendered
    assert "=B2*2" in rendered, "formulas must be preserved verbatim"
    assert "A4:B4" in rendered, "merged ranges must be reported"


def test_xlsx_sheet_target_read(xlsx_file):
    index = office.parse_document(xlsx_file)
    sheet = office.read_target(index, "说明")
    assert "附注" in sheet
    assert "小明" not in sheet


def test_xlsx_indexes_rows_beyond_400_and_supports_range_reads(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "大表"
    for row in range(1, 451):
        sheet.cell(row, 1).value = f"row-{row}"
    path = tmp_path / "large.xlsx"
    workbook.save(path)

    index = office.parse_document(path)
    assert index.blocks[0]["rows"] == 450
    assert "第 401 行" in index.blocks[0]["text"]
    window = office.read_target(index, "大表!A401:C405")
    assert "row-401" in window and "row-405" in window
    assert "row-400" not in window and "row-406" not in window


def test_xlsx_ranges_use_physical_rows_when_blank_rows_are_present(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "稀疏"
    sheet["A1"] = "标题"
    sheet["C3"] = "row-3"
    sheet["B100"] = "row-100"
    sheet["D400"] = "row-400"
    sheet["A401"] = "row-401"
    sheet.row_dimensions[400].hidden = True
    path = tmp_path / "sparse.xlsx"
    workbook.save(path)

    index = office.parse_document(path)
    window = office.read_target(index, "稀疏!A398:D401")
    assert "第 398 行" in window and "第 399 行" in window
    assert "第 400 行 [隐藏]" in window and "row-400" in window
    assert "row-401" in window and "第 397 行" not in window


def test_xlsx_range_preserves_hidden_row_marker_when_slicing_columns(tmp_path):
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "隐藏"
    sheet.append(["A", "B", "C"])
    sheet.append(["keep-a", "keep-b", "drop-c"])
    sheet.row_dimensions[2].hidden = True
    path = tmp_path / "hidden.xlsx"
    workbook.save(path)

    index = office.parse_document(path)
    window = office.read_target(index, "隐藏!A2:B2")

    assert "第 2 行 [隐藏]: A2=keep-a | B2=keep-b" in window
    assert "drop-c" not in window


def _make_exact_range_workbook(path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    for row in range(1, 451):
        sheet.cell(row, 1).value = f"A-{row}"
        sheet.cell(row, 2).value = f"B-{row}"
        sheet.cell(row, 3).value = f"C-{row}"
        sheet.cell(row, 4).value = f"D-{row}"
    # Physical gaps must remain gaps in the target result.
    for row in (2, 100, 400):
        for column in range(1, 5):
            sheet.cell(row, column).value = None
    # Sparse single-column data and a formula with no calculated cache in
    # openpyxl's normal writer path.
    for row in range(10, 21):
        sheet.cell(row, 3).value = None
    sheet["C10"] = "=A10+B10"
    sheet["C20"] = "sparse-c20"
    sheet.row_dimensions[7].hidden = True
    sheet.merge_cells("A20:B20")
    sheet["A20"] = "merged-a20"
    workbook.save(path)


def test_xlsx_single_column_range_never_leaks_other_columns(tmp_path):
    path = tmp_path / "exact.xlsx"
    _make_exact_range_workbook(path)

    window = office.read_target(office.parse_document(path), "Sheet1!A401:A450")

    assert "A401=A-401" in window and "A450=A-450" in window
    assert "B401=" not in window and "C401=" not in window and "D401=" not in window
    assert "B-401" not in window and "C-401" not in window and "D-401" not in window


def test_xlsx_sparse_single_column_uses_physical_rows(tmp_path):
    path = tmp_path / "sparse-column.xlsx"
    _make_exact_range_workbook(path)

    window = office.read_target(office.parse_document(path), "Sheet1!C10:C20")

    assert "C10==A10+B10" in window
    assert "C20=sparse-c20" in window
    assert "C11=（空）" in window and "C19=（空）" in window
    assert "A10=" not in window and "B10=" not in window and "D10=" not in window


def test_xlsx_single_row_and_single_cell_are_exact(tmp_path):
    path = tmp_path / "row-cell.xlsx"
    _make_exact_range_workbook(path)
    index = office.parse_document(path)

    row = office.read_target(index, "Sheet1!A10:D10")
    cell = office.read_target(index, "Sheet1!B5")

    assert all(f"{column}10=" in row for column in "ABCD")
    assert "A9=" not in row and "A11=" not in row
    assert "B5=B-5" in cell
    assert "A5=" not in cell and "C5=" not in cell and "D5=" not in cell


def test_xlsx_target_column_beyond_used_range_is_not_clamped(tmp_path):
    """An out-of-range coordinate must not silently become the last column."""

    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    sheet.append(["A5", "B5", "C5", "D5"])
    path = tmp_path / "column-boundary.xlsx"
    workbook.save(path)

    result = office.read_target(office.parse_document(path), "Sheet1!Z5")

    assert "Z5=（超出工作表索引范围 A:D）" in result
    assert "D5=D5" not in result
    assert "A5=A5" not in result


def test_xlsx_empty_rows_hidden_merge_and_formula_stay_precise(tmp_path):
    path = tmp_path / "metadata.xlsx"
    _make_exact_range_workbook(path)
    index = office.parse_document(path)

    empty_rows = office.read_target(index, "Sheet1!A1:A3")
    hidden = office.read_target(index, "Sheet1!A7:B7")
    merged = office.read_target(index, "Sheet1!A20:B20")
    formula = office.read_target(index, "Sheet1!C10")

    assert "第 2 行: A2=（空）" in empty_rows
    assert "第 7 行 [隐藏]: A7=A-7 | B7=B-7" in hidden
    assert "A20=merged-a20" in merged and "B20=（合并单元格 A20:B20）" in merged
    assert "C10==A10+B10" in formula and "缓存值" in formula


def test_office_indexes_embedded_docx_image_bytes(tmp_path):
    import io

    import docx
    from PIL import Image

    image = io.BytesIO()
    Image.new("RGB", (12, 8), "red").save(image, format="PNG")
    document = docx.Document()
    document.add_paragraph("图文混排")
    document.add_picture(io.BytesIO(image.getvalue()))
    path = tmp_path / "visual.docx"
    document.save(path)

    index = office.parse_document(path)
    visuals = index.meta.get("images")
    assert visuals and visuals[0]["data"].startswith(b"\x89PNG")


def test_office_render_pages_produces_full_page_visual(monkeypatch, tmp_path, pptx_file):
    """The renderer sends a complete slide page, not only embedded pictures."""

    import fitz

    monkeypatch.setattr(office.shutil, "which", lambda name: "fake-soffice")

    def fake_run(command, **kwargs):
        output = Path(command[command.index("--outdir") + 1])
        output.mkdir(parents=True, exist_ok=True)
        document = fitz.open()
        document.new_page(width=400, height=300)
        page = document.new_page(width=400, height=300)
        page.insert_text((30, 60), "FULL SLIDE ARROW AND SHAPE LAYOUT")
        document.save(str(output / "a.pdf"))
        document.close()

    monkeypatch.setattr(office.subprocess, "run", fake_run)
    workdir = tmp_path / "request"
    workdir.mkdir()
    images = office.render_pages(pptx_file, workdir, target="2", max_pages=2)
    assert len(images) == 1
    assert Path(images[0]).name == "office-page-02.png"
    assert Path(images[0]).stat().st_size > 100


def test_hybrid_pdf_with_text_and_scan_image_is_marked_for_ocr(tmp_path):
    import fitz
    from PIL import Image, ImageDraw

    image_path = tmp_path / "scan.png"
    image = Image.new("RGB", (500, 160), "white")
    ImageDraw.Draw(image).text((20, 50), "UNIQUE SCANNED LABEL", fill="black")
    image.save(image_path)
    pdf_path = tmp_path / "hybrid.pdf"
    document = fitz.open()
    page = document.new_page(width=600, height=800)
    page.insert_textbox((30, 30, 560, 300), "normal text " * 80)
    page.insert_image(fitz.Rect(60, 350, 540, 510), filename=str(image_path))
    document.save(str(pdf_path))
    document.close()

    from decision.multimodal.runtime import MultimodalRuntime

    assert MultimodalRuntime._pdf_sparse_pages(pdf_path) == [0]
    workdir = tmp_path / "regions"
    workdir.mkdir()
    regions = MultimodalRuntime._pdf_image_regions(pdf_path, workdir, [0])
    assert regions and Path(regions[0]).exists()


def test_office_rejects_other_formats(tmp_path):
    path = tmp_path / "a.txt"
    path.write_text("x", encoding="utf-8")
    assert office.is_office(path) is False
    with pytest.raises(ValueError):
        office.parse_document(path)


# ---------------------------------------------------------------- local OCR
def test_ocr_status_reports_missing_backends_with_hints():
    status = ocr.engine_status()
    assert set(status) >= {"tesseract", "paddleocr", "any"}
    if not status["tesseract"]["available"]:
        assert "pip install pytesseract" in status["tesseract"]["reason"]


def test_resolve_engine_uses_requested_backend_when_available(monkeypatch):
    monkeypatch.setattr(
        ocr,
        "engine_status",
        lambda: {
            "tesseract": {"available": True, "reason": ""},
            "paddleocr": {"available": False, "reason": "未安装 paddleocr"},
            "any": True,
        },
    )
    assert ocr.resolve_engine("tesseract")[0] == "tesseract"
    # A missing request falls back to what is installed, with a note.
    chosen, note = ocr.resolve_engine("paddleocr")
    assert chosen == "tesseract" and "PaddleOCR不可用" in note


def test_resolve_engine_reports_install_hint_when_nothing_is_available(monkeypatch):
    monkeypatch.setattr(
        ocr,
        "engine_status",
        lambda: {
            "tesseract": {"available": False, "reason": "未安装 pytesseract"},
            "paddleocr": {"available": False, "reason": "未安装 paddleocr"},
            "any": False,
        },
    )
    chosen, note = ocr.resolve_engine("auto")
    assert chosen == "" and "未安装" in note


@pytest.mark.asyncio
async def test_ocr_with_no_backend_returns_error_not_exception(monkeypatch, tmp_path):
    monkeypatch.setattr(ocr, "resolve_engine", lambda requested: ("", "本地未安装 OCR 引擎"))
    result = await ocr.ocr_image(tmp_path / "x.png")
    assert not result.ok
    assert "未安装" in result.error


@pytest.mark.asyncio
async def test_tesseract_backend_is_actually_invoked(monkeypatch, tmp_path):
    from PIL import Image

    path = tmp_path / "x.png"
    Image.new("RGB", (10, 10), (255, 255, 255)).save(path)
    calls = {}

    def fake_backend(target, lang):
        calls["lang"] = lang
        return "第一行\n第二行\n"

    monkeypatch.setattr(ocr, "_ocr_tesseract", fake_backend)
    monkeypatch.setattr(ocr, "resolve_engine", lambda requested: ("tesseract", ""))
    result = await ocr.ocr_image(path, lang="chi_sim+eng")
    assert result.ok
    assert calls["lang"] == "chi_sim+eng"
    assert result.text == "第一行\n第二行"


@pytest.mark.asyncio
async def test_empty_ocr_output_is_reported_as_failure(monkeypatch, tmp_path):
    monkeypatch.setattr(ocr, "_ocr_tesseract", lambda target, lang: "   \n")
    monkeypatch.setattr(ocr, "resolve_engine", lambda requested: ("tesseract", ""))
    result = await ocr.ocr_image(tmp_path / "x.png")
    assert not result.ok and "没有识别出文字" in result.error


# ------------------------------------------------------------ native video
class GeminiProvider:
    def __init__(self, key="k-123", base="https://example.invalid"):
        self.provider_config = {"type": "gemini_chat_completion", "api_base": base}
        self._key = key

    def get_current_key(self):
        return self._key


class OpenAILikeProvider:
    provider_config = {"type": "openai_chat_completion", "api_base": "https://x.invalid"}

    def get_current_key(self):
        return "sk-test"


def test_gemini_credentials_are_reused_not_reconfigured():
    key, base = native_video.gemini_credentials(GeminiProvider())
    assert key == "k-123" and base == "https://example.invalid"


def test_non_gemini_provider_never_claims_native_video():
    ok, reason = native_video.native_available(OpenAILikeProvider())
    assert ok is False and "Gemini" in reason


def test_missing_key_disables_native_video():
    ok, reason = native_video.native_available(GeminiProvider(key=""))
    assert ok is False


@pytest.mark.asyncio
async def test_ocr_fallback_note_does_not_discard_recognized_text(monkeypatch, tmp_path):
    from PIL import Image

    image = tmp_path / "ocr.png"
    Image.new("RGB", (16, 16), "white").save(image)
    monkeypatch.setattr(
        ocr,
        "resolve_engine",
        lambda requested: ("tesseract", "已从 PaddleOCR 降级"),
    )
    monkeypatch.setattr(
        ocr, "_ocr_tesseract", lambda path, lang: "降级后仍识别到的文字"
    )

    result = await ocr.ocr_image(image, engine="paddleocr")
    assert result.ok is True
    assert result.text == "降级后仍识别到的文字"
    assert result.error == "已从 PaddleOCR 降级"


def test_capability_label_requires_a_real_native_route():
    from decision.multimodal import adapters

    gemini = GeminiProvider()
    assert adapters.classify_kind(gemini, "video") == "native"
    frames_only = SimpleNamespace(
        provider_config={"modalities": ["text", "image"], "type": "openai_chat_completion"},
        get_current_key=lambda: "",
    )
    assert adapters.classify_kind(frames_only, "video") == "preprocessed"


def test_interactions_parser_excludes_user_thought_and_tools():
    payload = {
        "status": "completed",
        "steps": [
            {"type": "user_input", "content": [{"type": "text", "text": "秘密输入"}]},
            {"type": "thought", "summary": [{"type": "text", "text": "内部思考"}]},
            {"type": "function_call", "arguments": {"text": "工具参数"}},
            {"type": "function_result", "content": [{"type": "text", "text": "工具结果"}]},
            {"type": "model_output", "content": [{"type": "text", "text": "最终视频答案"}]},
        ],
    }
    result = native_video._model_output_text(payload)
    assert result == "最终视频答案"
    assert all(value not in result for value in ("秘密输入", "内部思考", "工具结果"))


@pytest.mark.asyncio
async def test_native_upload_failure_degrades_to_frames(monkeypatch, tmp_path):
    async def boom(*args, **kwargs):
        raise RuntimeError("upload refused")

    monkeypatch.setattr(native_video, "upload_file", boom)
    path = tmp_path / "v.mp4"
    path.write_bytes(b"x")
    result = await native_video.analyze(
        path, "describe", provider=GeminiProvider(), model="gemini-2.5-flash"
    )
    assert result.route == "frames"
    assert "upload refused" in result.error


@pytest.mark.asyncio
async def test_native_success_returns_native_route_and_cleans_up(monkeypatch, tmp_path):
    deleted = []
    uploaded = {}

    async def fake_upload(path, *, api_key, base, timeout=300.0):
        uploaded["key"] = api_key
        return "files/abc", "https://files/abc"

    async def fake_wait(name, **kwargs):
        uploaded["waited"] = name

    async def fake_generate(**kwargs):
        uploaded["uri"] = kwargs["file_uri"]
        return "视频里有一只猫"

    async def fake_delete(name, **kwargs):
        deleted.append(name)

    monkeypatch.setattr(native_video, "upload_file", fake_upload)
    monkeypatch.setattr(native_video, "wait_active", fake_wait)
    monkeypatch.setattr(native_video, "generate", fake_generate)
    monkeypatch.setattr(native_video, "delete_file", fake_delete)
    path = tmp_path / "v.mp4"
    path.write_bytes(b"x")
    result = await native_video.analyze(
        path, "describe", provider=GeminiProvider(), model="gemini-2.5-flash"
    )
    assert result.route == "native" and "猫" in result.text
    assert uploaded["uri"] == "https://files/abc"
    assert deleted == ["files/abc"], "the uploaded file must be removed afterwards"


@pytest.mark.asyncio
async def test_native_cancellation_still_cleans_up(monkeypatch, tmp_path):
    deleted = []

    async def fake_upload(path, **kwargs):
        return "files/x", "https://files/x"

    async def fake_wait(name, **kwargs):
        return None

    async def fake_generate(**kwargs):
        await asyncio.sleep(30)
        return "never"

    async def fake_delete(name, **kwargs):
        deleted.append(name)

    monkeypatch.setattr(native_video, "upload_file", fake_upload)
    monkeypatch.setattr(native_video, "wait_active", fake_wait)
    monkeypatch.setattr(native_video, "generate", fake_generate)
    monkeypatch.setattr(native_video, "delete_file", fake_delete)
    path = tmp_path / "v.mp4"
    path.write_bytes(b"x")
    task = asyncio.create_task(
        native_video.analyze(path, "p", provider=GeminiProvider(), model="m")
    )
    await asyncio.sleep(0.05)
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert deleted == ["files/x"]


@pytest.mark.asyncio
async def test_native_agentic_video_uses_interactions_endpoint(monkeypatch, tmp_path):
    calls = {}

    async def fake_upload(path, **kwargs):
        return "files/agentic", "https://files/agentic"

    async def fake_wait(name, **kwargs):
        return None

    async def fake_agentic(**kwargs):
        calls.update(kwargs)
        return "长视频 Agentic 结果"

    async def fake_delete(name, **kwargs):
        calls["deleted"] = name

    monkeypatch.setattr(native_video, "upload_file", fake_upload)
    monkeypatch.setattr(native_video, "wait_active", fake_wait)
    monkeypatch.setattr(native_video, "agentic_generate", fake_agentic)
    monkeypatch.setattr(native_video, "delete_file", fake_delete)
    path = tmp_path / "long.mp4"
    path.write_bytes(b"x")

    result = await native_video.analyze(
        path,
        "精读",
        provider=GeminiProvider(),
        model="gemini-2.5-flash",
        agentic=True,
    )
    assert result.route == "native" and result.agentic is True
    assert calls["file_uri"] == "https://files/agentic"
    assert calls["deleted"] == "files/agentic"
